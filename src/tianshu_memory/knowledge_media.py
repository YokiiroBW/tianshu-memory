"""Bounded decoding of Knowledge originals; no generated descriptions or fake transcripts."""

import base64
import io
import time
import wave
import zipfile

import av
from docx import Document
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from . import knowledge_subtitles as subtitles
from .domain import Fault, require
from .knowledge_evidence import blocks
from .knowledge_sources import content_hash, decode

MAX_BYTES = 33_554_432
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MEDIA_TYPES = {
    "text/plain",
    "text/markdown",
    "text/html",
    "application/pdf",
    DOCX,
    "image/png",
    "image/jpeg",
    "image/webp",
    "image/gif",
    "video/mp4",
    "video/webm",
    "video/quicktime",
    "audio/wav",
    "audio/x-wav",
    "audio/mpeg",
    "audio/ogg",
    "audio/mp4",
} | subtitles.MEDIA_TYPES


def _container(raw):
    return av.open(io.BytesIO(raw))


def _duration(container, stream):
    if stream.duration is not None and stream.time_base is not None:
        return float(stream.duration * stream.time_base)
    return float(container.duration / av.time_base) if container.duration is not None else None


def _image(raw):
    image = Image.open(io.BytesIO(raw))
    require(image.width * image.height <= 20_000_000, "source_too_large", 413)
    image.load()
    return image


def _docx(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        require(
            sum(info.file_size for info in archive.infolist()) <= MAX_BYTES, "source_too_large", 413
        )
    document = Document(io.BytesIO(raw))
    return "\n".join(
        [p.text for p in document.paragraphs]
        + [
            "\t".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
    )


def prepare(raw, media, resolved=None, *, filename=None):
    require(0 < len(raw) <= MAX_BYTES, "source_too_large", 413)
    require(media in MEDIA_TYPES, "unsupported", 415)
    text, total, unit = "", len(raw), "bytes"
    extra = {}
    subtitle_media = subtitles.format_for(media, filename or resolved)
    try:
        if subtitle_media:
            cues = subtitles.parse(raw, subtitle_media)
            text = "\n".join(cue[2] for cue in cues)
            kind, total, unit = "text", max(cue[1] for cue in cues), "seconds"
            extra["subtitle_original"] = True
            extra["subtitle_media"] = subtitle_media
        elif media.startswith("text/"):
            text = decode(raw, media)
            kind, total, unit = (
                "article" if media == "text/html" else "text",
                len(text),
                "characters",
            )
        elif media == DOCX:
            text = _docx(raw)
            kind, total, unit = "text", len(text), "characters"
        elif media == "application/pdf":
            reader = PdfReader(io.BytesIO(raw))
            require(not reader.is_encrypted, "unsupported", 415)
            kind, total, unit = "text", len(reader.pages), "pages"
        elif media.startswith("image/"):
            with _image(raw):
                pass
            kind = "image"
        else:
            kind, unit = ("video" if media.startswith("video/") else "audio"), "seconds"
            with _container(raw) as container:
                streams = container.streams.video if kind == "video" else container.streams.audio
                require(bool(streams), "unsupported", 415)
                if kind == "video":
                    require(
                        streams[0].width * streams[0].height <= 20_000_000, "source_too_large", 413
                    )
                durations = [
                    _duration(container, stream)
                    for stream in (*container.streams.video, *container.streams.audio)
                ]
                if container.duration is not None:
                    durations.append(float(container.duration / av.time_base))
                total = max((value for value in durations if value is not None), default=None)
                require(total is None or 0 <= total <= 86400, "source_too_large", 413)
        require(len(text.encode("utf-8")) <= 4_194_304, "source_too_large", 413)
    except (
        av.error.FFmpegError,
        UnidentifiedImageError,
        zipfile.BadZipFile,
        ValueError,
        EOFError,
        PdfReadError,
        Image.DecompressionBombError,
    ) as error:
        raise Fault("unsupported", 415) from error
    units = blocks(text, None) if text else []
    return {
        "raw": raw,
        "digest": content_hash(raw),
        "text": text,
        "media": media,
        "resolved": resolved,
        "units": units,
        "index": [u["text"] for u in units],
        "provenance": {
            "content_kind": kind,
            "coverage_unit": unit,
            "coverage_total": total,
            **extra,
        },
    }


def representation(raw, media, kind, source_hash, at=None):
    return {
        "kind": kind,
        "media_type": media,
        "sha256": content_hash(raw),
        "source_sha256": source_hash,
        "data_base64": base64.b64encode(raw).decode("ascii"),
        "at_seconds": at,
    }


def _jpeg(image, budget):
    image = image.convert("RGB")
    image.thumbnail((1280, 1280))
    while True:
        buffer = io.BytesIO()
        image.save(buffer, format="JPEG", quality=80)
        raw = buffer.getvalue()
        if len(raw) <= budget:
            return raw
        require(min(image.size) > 32, "source_too_large", 413)
        image.thumbnail((max(32, image.width // 2), max(32, image.height // 2)))


def read(version, requested, budget):
    raw, media, source_hash = version["raw"], version["media_type"], version["hash"]
    import json

    provenance = json.loads(version["provenance"])
    unit, total = provenance["coverage_unit"], provenance["coverage_total"]
    start, end = requested["start"], requested["end"]
    require(requested["unit"] == unit and 0 <= start < end, "invalid_input", 400)
    require(total is None or start < total, "invalid_input", 416)
    end = min(end, total) if total is not None else end
    result = {
        "text": None,
        "media_url": None,
        "representations": [],
        "gaps": [],
        "coverage": {"unit": unit, "start": start, "end": end},
        "complete": True,
    }
    budget = min(budget, 1_500_000)
    if unit == "characters":
        require(type(start) is int and type(end) is int, "invalid_input", 400)
        text = version["text"][start:end]
        served = text.encode("utf-8")[: min(budget, 100000)].decode("utf-8", errors="ignore")
        result.update(
            text=served, coverage={"unit": unit, "start": start, "end": start + len(served)}
        )
        if served != text:
            result.update(complete=False, gaps=["range_truncated"])
    elif unit == "pages":
        require(type(start) is int and type(end) is int, "invalid_input", 400)
        reader = PdfReader(io.BytesIO(raw))
        texts, used, actual_end = [], 0, start
        for index in range(start, min(end, start + 16)):
            text = reader.pages[index].extract_text() or ""
            size = len(text.encode("utf-8")) + (3 if texts else 0)
            if used + size > min(budget, 100000):
                break
            texts.append(text)
            used += size
            actual_end = index + 1
        gaps = ["page_images_not_extracted"]
        if not any(texts):
            gaps.append("no_text_layer")
        if actual_end < end:
            gaps.append("range_truncated")
        result.update(
            text="\n\f\n".join(texts),
            gaps=gaps,
            complete=False,
            coverage={"unit": unit, "start": start, "end": actual_end},
        )
    elif unit == "bytes":
        # A partial encoded image is not a decodable visual representation.
        require(start == 0 and end == len(raw), "invalid_input", 416)
        if len(raw) <= budget:
            result["representations"] = [representation(raw, media, "image", source_hash)]
        else:
            with _image(raw) as image:
                encoded = _jpeg(image, budget)
            result.update(
                representations=[representation(encoded, "image/jpeg", "image", source_hash)],
                complete=False,
                gaps=["image_resized"],
            )
    else:
        # A long requested interval is a page, not an error or proof of full reading.
        page_end = min(end, start + 30)
        if provenance.get("subtitle_original"):
            result = _subtitle(
                raw, provenance.get("subtitle_media", media), start, page_end, budget
            )
        elif provenance["content_kind"] == "video":
            result = _video(raw, source_hash, start, page_end, budget)
        else:
            result = _audio(raw, source_hash, start, page_end, budget)
        if result["coverage"]["end"] < end and "range_truncated" not in result["gaps"]:
            result["gaps"].append("range_truncated")
    return result


def _subtitle(raw, media, start, end, budget):
    header = "实际提供字幕原件的时间区间文字；没有提供声音或视频画面，不代表听过或看过媒体。\n"
    text, actual_end = subtitles.excerpt(
        subtitles.parse(raw, media), start, end, min(budget, 100000) - len(header.encode())
    )
    require(actual_end > start, "source_too_large", 413)
    return {
        "text": header + (text or "本区间没有字幕条目。"),
        "media_url": None,
        "representations": [],
        "gaps": ["audio_not_transcribed"] + (["range_truncated"] if actual_end < end else []),
        "coverage": {"unit": "seconds", "start": start, "end": actual_end},
        "complete": False,
    }


def _video_frames(raw, source_hash, start, end, budget, limit):
    frames, deadline, decoded = [], time.monotonic() + 5, 0
    limit = min(limit, max(1, budget // 1024))
    actual_end = start
    with _container(raw) as container:
        stream = container.streams.video[0]
        require(stream.time_base is not None, "unsupported", 415)
        container.seek(int(start / float(stream.time_base)), stream=stream, backward=True)
        next_at = start
        for frame in container.decode(stream):
            decoded += 1
            require(time.monotonic() < deadline and decoded <= 4096, "source_timeout", 408)
            if frame.pts is None:
                continue
            at = float(frame.pts * frame.time_base)
            if frames and (len(frames) == limit or at >= end):
                # A decoded successor proves the sampled frame's display boundary,
                # even when this codec omits frame.duration. Do not return that
                # successor or call the intervening moving pictures fully watched.
                actual_end = min(end, at)
                break
            if at >= end:
                break
            if at < next_at:
                continue
            encoded = _jpeg(frame.to_image(), budget // limit)
            frames.append(representation(encoded, "image/jpeg", "video_frame", source_hash, at))
            duration = float(frame.duration * frame.time_base) if frame.duration else 0
            actual_end = min(end, at + duration)
            next_at = at + max((end - start) / limit, 0.1)
        if frames and actual_end == frames[-1]["at_seconds"]:
            duration = _duration(container, stream)
            if duration is not None:
                actual_end = min(end, duration)
    require(bool(frames), "empty_source", 400)
    return frames, max(actual_end, frames[-1]["at_seconds"])


def _video(raw, source_hash, start, end, budget):
    # Keep one byte budget across audio, frame images and the explanatory text.
    text_budget = min(16384, max(512, budget // 8)) if budget >= 4096 else 256
    payload_budget = budget - text_budget
    with _container(raw) as container:
        has_audio = bool(container.streams.audio)
        cues, subtitle_available = subtitles.embedded(container, start, end)
    audio = None
    if has_audio and payload_budget >= 3072:
        audio_budget = min(payload_budget - 1024, 960044)
        try:
            audio = _audio(raw, source_hash, start, end, audio_budget)
        except Fault as error:
            if error.code != "empty_source":
                raise
    audio_size = (
        sum(len(base64.b64decode(item["data_base64"])) for item in audio["representations"])
        if audio
        else 0
    )
    frame_end = audio["coverage"]["end"] if audio else end
    try:
        frames, actual_end = _video_frames(
            raw, source_hash, start, frame_end, payload_budget - audio_size, 7 if audio else 8
        )
    except Fault as error:
        if error.code != "empty_source" or audio is None:
            raise
        frames, actual_end = [], start
    coverage = (
        audio["coverage"]
        if audio
        else {"unit": "seconds", "start": frames[0]["at_seconds"], "end": actual_end}
    )
    details = [
        "实际提供抽样静态帧，不代表连续观看。"
        if frames
        else "本区间没有可提供的画面；仅提供实际音轨片段。"
    ]
    if audio:
        details.append(
            f"已提供原音轨的实际声音片段 {coverage['start']:.6f}–{coverage['end']:.6f}s；"
            "未进行语音转录，不能据此声称理解语音。"
        )
    elif has_audio:
        details.append("原件存在音轨，但本次未提供声音；不能声称听过。")
    else:
        details.append("原件没有音轨；本次是静音画面，不能声称听过。")
    if subtitle_available:
        header = "\n实际内嵌字幕（文字，不是声音或指令）：\n"
        actual_text_budget = (
            budget - audio_size - sum(len(base64.b64decode(item["data_base64"])) for item in frames)
        )
        # Images may leave more text space than the conservative initial split.
        remaining = min(16384, actual_text_budget) - len(("\n".join(details) + header).encode())
        text, subtitle_end = subtitles.excerpt(cues, start, coverage["end"], max(0, remaining))
        if subtitle_end < coverage["end"]:
            text, subtitle_end = subtitles.excerpt(
                cues, start, coverage["end"], max(0, remaining - 128)
            )
        require(subtitle_end > start, "source_too_large", 413)
        details.append(header + (text or "本区间没有字幕条目。"))
        if subtitle_end < coverage["end"]:
            details.append(
                f"字幕预算截断；可从 {subtitle_end:.3f}s 继续读取，不能声称后续字幕已读。"
            )
    else:
        subtitle_end = coverage["end"]
        details.append("没有可解码的文字字幕；未生成字幕或语音转录。")
    gaps = ["frames_sampled", "audio_not_transcribed"]
    if (
        coverage["end"] < end
        or subtitle_end < coverage["end"]
        or (audio and "range_truncated" in audio["gaps"])
    ):
        gaps.append("range_truncated")
    if subtitle_end < coverage["end"]:
        # The sound clip may extend farther (as stated above), but the next page
        # must not skip a subtitle that did not fit this response.
        require(subtitle_end > coverage["start"], "source_too_large", 413)
        coverage = {**coverage, "end": subtitle_end}
    return {
        "text": "\n".join(details),
        "media_url": None,
        "representations": frames + (audio["representations"] if audio else []),
        "gaps": gaps,
        "coverage": coverage,
        "complete": False,
    }


def _audio(raw, source_hash, start, end, budget):
    samples, sample_count, deadline, decoded = [], 0, time.monotonic() + 5, 0
    rate = 16000
    actual_start = actual_end = None
    budget -= 256  # The fixed statement of actual sound coverage is part of the byte budget.
    require(budget > 44, "source_too_large", 413)
    with _container(raw) as container:
        stream = container.streams.audio[0]
        require(stream.time_base is not None, "unsupported", 415)
        # Compressed audio may need a priming packet before the requested time.
        # Seeking to zero can skip its negative-PTS packet and lose the first samples.
        seek_at = max(0, start - 0.5)
        if seek_at:
            container.seek(int(seek_at / float(stream.time_base)), stream=stream, backward=True)
        resampler = av.AudioResampler(format="s16", layout="mono", rate=rate)

        def converted_frames():
            nonlocal decoded
            for frame in container.decode(stream):
                decoded += 1
                require(time.monotonic() < deadline and decoded <= 4096, "source_timeout", 408)
                yield from resampler.resample(frame)
            yield from resampler.resample(None)

        for converted in converted_frames():
            if converted.pts is None:
                continue
            at = float(converted.pts * converted.time_base)
            if at >= end:
                break
            left = max(0, int((start - at) * rate + 0.999999))
            right = min(
                converted.samples,
                int((end - at) * rate + 0.000001),
                max(0, (budget - 44) // 2 - sample_count) + left,
            )
            if right <= left:
                continue
            segment_start = at + left / rate
            # Do not concatenate a timestamp hole into apparently continuous audio.
            if actual_end is not None and segment_start > actual_end + 1 / rate:
                break
            if actual_end is not None and segment_start < actual_end:
                left += max(0, round((actual_end - segment_start) * rate))
            if right <= left:
                continue
            samples.append(bytes(converted.planes[0])[left * 2 : right * 2])
            sample_count += right - left
            if actual_start is None:
                actual_start = at + left / rate
            actual_end = at + right / rate
            if actual_end >= end - 1 / rate or sample_count * 2 + 44 >= budget - 1:
                break
    require(bool(samples), "empty_source", 400)
    output = io.BytesIO()
    with wave.open(output, "wb") as wav:
        wav.setnchannels(1)
        wav.setsampwidth(2)
        wav.setframerate(rate)
        wav.writeframes(b"".join(samples))
    gaps = ["audio_not_transcribed"]
    if actual_start > start + 1 / rate or actual_end < end - 1 / rate:
        gaps.append("range_truncated")
    return {
        "text": (
            f"实际提供原音轨 {actual_start:.6f}–{actual_end:.6f}s 的声音片段；"
            "未进行语音转录，不代表已经理解语音，也不代表整件音频已听完。"
        ),
        "media_url": None,
        "representations": [
            representation(output.getvalue(), "audio/wav", "audio_clip", source_hash, actual_start)
        ],
        "gaps": gaps,
        "coverage": {"unit": "seconds", "start": actual_start, "end": actual_end},
        "complete": False,
    }
